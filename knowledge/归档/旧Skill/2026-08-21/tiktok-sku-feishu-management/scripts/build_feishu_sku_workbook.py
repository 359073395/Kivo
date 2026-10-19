import json
import os
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


DATA = Path(os.environ["DATA_JSON"])
OUT = Path(os.environ["OUT_XLSX"])

payload = json.loads(DATA.read_text(encoding="utf-8"))

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)
thin_gray = Side(style="thin", color="D9E2F3")
border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
note_fill = PatternFill("solid", fgColor="FFF2CC")
normal_fill = PatternFill("solid", fgColor="E2F0D9")
presale_fill = PatternFill("solid", fgColor="FCE4D6")
remove_fill = PatternFill("solid", fgColor="F4CCCC")
unknown_fill = PatternFill("solid", fgColor="E7E6E6")


def add_sheet(name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    add_table(ws, name)
    return ws


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border


def add_table(ws, base_name):
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table_name = "".join(ch for ch in base_name if ch.isalnum())[:20] or "Table"
    table = Table(displayName=f"{table_name}Table", ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def autosize(ws, caps=None):
    caps = caps or {}
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            text = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(text), 80))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), caps.get(letter, 45))


def add_status_validation(ws, col_letter):
    if isinstance(col_letter, int):
        col_letter = get_column_letter(col_letter)
    dv = DataValidation(type="list", formula1='"正常,下架,开预购,关预购,待确认"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}5000")


def add_status_colors(ws, status_col):
    last_row = max(ws.max_row, 5000)
    status_letter = get_column_letter(status_col)
    full_range = f"A2:{get_column_letter(ws.max_column)}{last_row}"
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(
            formula=[f'${status_letter}2="下架"'],
            fill=remove_fill,
            font=Font(color="9C0006", bold=True),
        ),
    )
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(
            formula=[f'${status_letter}2="开预购"'],
            fill=presale_fill,
            font=Font(color="C65911", bold=True),
        ),
    )
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(
            formula=[f'${status_letter}2="关预购"'],
            fill=normal_fill,
            font=Font(color="006100", bold=True),
        ),
    )
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(
            formula=[f'${status_letter}2="正常"'],
            fill=normal_fill,
            font=Font(color="375623"),
        ),
    )
    ws.conditional_formatting.add(
        full_range,
        FormulaRule(
            formula=[f'${status_letter}2="待确认"'],
            fill=unknown_fill,
            font=Font(color="666666"),
        ),
    )


today = date.today().isoformat()

guide = wb.create_sheet("使用说明")
guide_rows = [
    ["用途", "这是一份可导入飞书的 SKU 总控表。沟通和维护时以 SKU 为唯一依据。"],
    ["默认规则", "同一个 SKU 出现在多个店铺时，更新该 SKU 会影响所有匹配店铺。"],
    ["状态颜色", "下架=红色；开预购=橙色；关预购=浅绿色；正常=绿色；待确认=灰色。"],
    ["产品总表", "一行代表一个店铺里的一个 SKU，用于执行和检查每个店铺。"],
    ["SKU汇总", "一行代表一个 SKU，用于查看这个 SKU 覆盖了哪些店铺。"],
    ["变更记录", "以后每次下架、开预购、关预购、恢复正常都可以记录在这里。"],
    ["异常记录", "包括空 SKU、同一店铺内重复 SKU、读取异常等。"],
    ["生成日期", today],
]
guide.append(["项目", "说明"])
for row in guide_rows:
    guide.append(row)
style_sheet(guide)
autosize(guide, {"A": 18, "B": 90})

product_headers = ["SKU", "店铺", "处理状态", "标注备注", "最近更新时间", "源文件", "原始行号"]
product_rows = [
    [
        r["sku"],
        r["store"],
        "",
        "",
        "",
        r["source_file"],
        r["source_row"],
    ]
    for r in sorted(payload["records"], key=lambda x: (x["sku"], x["store"], x["source_file"], x["source_row"]))
]
product_ws = add_sheet("产品总表", product_headers, product_rows)
add_status_validation(product_ws, 3)
add_status_colors(product_ws, 3)
autosize(product_ws, {"A": 32, "B": 24, "C": 14, "D": 40, "E": 16, "F": 30, "G": 12})

summary_headers = ["SKU", "出现店铺数", "涉及店铺", "总记录数", "处理状态", "标注备注", "最近更新时间"]
summary_rows = [
    [
        r["sku"],
        r["store_count"],
        r["stores"],
        r["row_count"],
        "",
        "",
        "",
    ]
    for r in payload["summary"]
]
summary_ws = add_sheet("SKU汇总", summary_headers, summary_rows)
add_status_validation(summary_ws, 5)
add_status_colors(summary_ws, 5)
autosize(summary_ws, {"A": 32, "B": 14, "C": 90, "D": 12, "E": 14, "F": 40, "G": 16})

stats_headers = ["店铺", "源文件", "原始行数", "唯一SKU数"]
stats_rows = [[r["store"], r["file"], r["rows"], r["unique_skus"]] for r in payload["file_stats"]]
stats_ws = add_sheet("店铺统计", stats_headers, stats_rows)
autosize(stats_ws, {"A": 24, "B": 32, "C": 12, "D": 12})

log_headers = ["时间", "SKU", "操作", "影响店铺数", "结果", "备注"]
log_rows = [["", "", "", "", "", ""] for _ in range(20)]
log_ws = add_sheet("变更记录", log_headers, log_rows)
dv_action = DataValidation(type="list", formula1='"下架,开预购,关预购,恢复正常,待确认"', allow_blank=True)
log_ws.add_data_validation(dv_action)
dv_action.add("C2:C5000")
autosize(log_ws, {"A": 18, "B": 32, "C": 14, "D": 14, "E": 24, "F": 50})

exception_headers = ["店铺", "源文件", "原始行号", "问题", "内容"]
exception_rows = [
    [r["store"], r["file"], r["row"], r["issue"], r["value"]]
    for r in payload["exceptions"]
]
exception_ws = add_sheet("异常记录", exception_headers, exception_rows)
autosize(exception_ws, {"A": 24, "B": 32, "C": 12, "D": 24, "E": 50})
for row in exception_ws.iter_rows(min_row=2):
    for cell in row:
        cell.fill = note_fill

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 24

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
