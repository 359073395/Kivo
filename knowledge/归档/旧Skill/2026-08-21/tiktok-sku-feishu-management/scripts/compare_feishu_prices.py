import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from shutil import which

from openpyxl import load_workbook


TOKEN = os.environ["SPREADSHEET_TOKEN"]
DATA = Path(os.environ["DATA_JSON"])
PRICE_XLSX = Path(os.environ["PRICE_XLSX"])
REPORT_JSON = Path(os.environ.get("REPORT_JSON", "price_compare_report.json"))


def cli_prefix():
    direct = os.environ.get("LARK_CLI_RUNJS")
    node = os.environ.get("NODE_EXE") or which("node.exe") or which("node")
    if direct and node:
        return [node, direct]
    command = os.environ.get("LARK_CLI") or which("lark-cli.cmd") or which("lark-cli")
    if command and node:
        runjs = Path(command).resolve().parent / "node_modules" / "@larksuite" / "cli" / "scripts" / "run.js"
        if runjs.exists():
            return [node, str(runjs)]
    if command:
        return [command]
    raise SystemExit("lark-cli is not installed or is not available on PATH.")


CLI_PREFIX = cli_prefix()

DISPATCH_TITLE = "\u5458\u5de5\u6267\u884c\u6e05\u5355"


def run_cli(args):
    proc = subprocess.run(
        [*CLI_PREFIX, *args],
        text=True,
        encoding="utf-8",
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if proc.returncode != 0:
        print(proc.stdout)
        print(proc.stderr, file=sys.stderr)
        raise SystemExit(f"Command failed: {' '.join(args)}")
    return json.loads(proc.stdout)


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return " ".join(text.split())


def number(value):
    text = clean(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def price_text(value):
    parsed = number(value)
    if parsed is None:
        return clean(value)
    if parsed == parsed.to_integral_value():
        return str(int(parsed))
    return format(parsed.normalize(), "f")


def price_for_sku(sku, prices):
    key = clean(sku).lower()
    if key in prices:
        return prices[key], "exact"
    match = re.match(r"^(\d+)\s*pcs\s*(.+)$", clean(sku), flags=re.IGNORECASE)
    if not match or match.group(2).lower() not in prices:
        return ["", "", ""], "missing"
    multiplier = int(match.group(1))
    base = prices[match.group(2).lower()]
    return [price_text(number(value) * multiplier) if number(value) is not None else "" for value in base], f"{multiplier}pcs"


def load_prices():
    sheet = load_workbook(PRICE_XLSX, read_only=True, data_only=True).worksheets[0]
    headers = None
    header_row = 0
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        values = [clean(value) for value in row]
        if "\u5546\u54c1 SKU" in values or "SKU" in values:
            headers, header_row = values, row_number
            break
    if headers is None:
        raise SystemExit("Cannot find SKU header in price workbook.")
    sku = headers.index("\u5546\u54c1 SKU") if "\u5546\u54c1 SKU" in headers else headers.index("SKU")
    original = headers.index("\u4e0a\u67b6\u539f\u4ef7")
    daily = headers.index("\u65e5\u5e38\u552e\u4ef7")
    activity = headers.index("\u5e73\u53f0\u6d3b\u52a8\u4ef7")
    result = {}
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        item_sku = clean(row[sku] if sku < len(row) else "")
        if item_sku:
            result[item_sku.lower()] = [price_text(row[original]), price_text(row[daily]), price_text(row[activity])]
    return result


def read_range(sheet_id, area):
    return run_cli(
        [
            "sheets",
            "+read",
            "--as",
            "user",
            "--spreadsheet-token",
            TOKEN,
            "--range",
            f"{sheet_id}!{area}",
            "--value-render-option",
            "FormattedValue",
        ]
    )["data"]["valueRange"].get("values", [])


records = sorted(json.loads(DATA.read_text(encoding="utf-8"))["records"], key=lambda r: (r["store"], r["sku"], r["source_row"]))
prices = load_prices()
info = run_cli(["sheets", "+info", "--as", "user", "--spreadsheet-token", TOKEN])
sheet_ids = {sheet["title"]: sheet["sheet_id"] for sheet in info["data"]["sheets"]["sheets"]}
match_kinds = defaultdict(int)
mismatches = []
missing = []

expected = []
for record in records:
    values, kind = price_for_sku(record["sku"], prices)
    match_kinds[kind] += 1
    expected.append(values)
    if not any(values):
        missing.append({"store": record["store"], "sku": record["sku"]})

dispatch = read_range(sheet_ids[DISPATCH_TITLE], f"H2:J{len(records) + 1}")
for index, (record, wanted) in enumerate(zip(records, expected), 2):
    got = dispatch[index - 2] if index - 2 < len(dispatch) else []
    got = [price_text(got[i]) if i < len(got) else "" for i in range(3)]
    if got != wanted:
        mismatches.append({"sheet": DISPATCH_TITLE, "row": index, "store": record["store"], "sku": record["sku"], "expected": wanted, "actual": got})

by_store = defaultdict(list)
for record in records:
    by_store[record["store"]].append(record)
for store, store_records in by_store.items():
    rows = sorted(store_records, key=lambda r: (r["sku"], r["source_row"]))
    values = read_range(sheet_ids[f"\u5e97\u94fa-{store}"[:100]], f"A2:H{len(rows) + 1}")
    for index, record in enumerate(rows, 2):
        wanted, _ = price_for_sku(record["sku"], prices)
        row = values[index - 2] if index - 2 < len(values) else []
        got = [price_text(row[i]) if i < len(row) else "" for i in range(5, 8)]
        if got != wanted:
            mismatches.append({"sheet": f"\u5e97\u94fa-{store}", "row": index, "sku": record["sku"], "expected": wanted, "actual": got})

report = {
    "price_file": str(PRICE_XLSX),
    "price_skus": len(prices),
    "records_checked": len(records),
    "store_sheets_checked": len(by_store),
    "match_kinds": dict(match_kinds),
    "missing_price_rows": len(missing),
    "mismatch_count": len(mismatches),
    "sample_missing": missing[:20],
    "sample_mismatches": mismatches[:20],
}
REPORT_JSON.parent.mkdir(parents=True, exist_ok=True)
REPORT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False))
