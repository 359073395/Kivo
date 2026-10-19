import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from shutil import which

from openpyxl import load_workbook


TOKEN = os.environ["SPREADSHEET_TOKEN"]
DATA = Path(os.environ["DATA_JSON"])
PRICE_XLSX = Path(os.environ["PRICE_XLSX"])


def cli_prefix():
    direct = os.environ.get("LARK_CLI_RUNJS")
    node = os.environ.get("NODE_EXE") or which("node.exe") or which("node")
    if direct and node:
        return [node, direct]
    command = os.environ.get("LARK_CLI") or which("lark-cli.cmd") or which("lark-cli")
    if command and node:
        runjs = Path(command).resolve().parent / "node_modules" / "@larksuite" / "cli" / "scripts" / "run.js"
        if runjs.exists():
            return [node, str(runjs)]
    if command:
        return [command]
    raise SystemExit("lark-cli is not installed or is not available on PATH.")


CLI_PREFIX = cli_prefix()


def run_cli(args):
    proc = subprocess.run(
        [*CLI_PREFIX, *args],
        text=True,
        encoding="utf-8",
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if proc.returncode != 0:
        print(proc.stdout)
        print(proc.stderr, file=sys.stderr)
        raise SystemExit(f"Command failed: {' '.join(args)}")
    return json.loads(proc.stdout)


def write_range(sheet_id, start_cell, values, chunk_size=20):
    col = "".join(ch for ch in start_cell if ch.isalpha())
    row = int("".join(ch for ch in start_cell if ch.isdigit()))

    def col_num(letter):
        out = 0
        for ch in letter:
            out = out * 26 + ord(ch.upper()) - 64
        return out

    def col_letter(num):
        out = ""
        while num:
            num, rem = divmod(num - 1, 26)
            out = chr(65 + rem) + out
        return out

    width = max(len(r) for r in values)
    c1 = col_num(col)
    for i in range(0, len(values), chunk_size):
        chunk = values[i : i + chunk_size]
        r1 = row + i
        r2 = r1 + len(chunk) - 1
        c2 = col_letter(c1 + width - 1)
        run_cli(
            [
                "sheets",
                "+write",
                "--as",
                "user",
                "--spreadsheet-token",
                TOKEN,
                "--range",
                f"{sheet_id}!{col}{r1}:{c2}{r2}",
                "--values",
                json.dumps(chunk, ensure_ascii=True),
            ]
        )


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return " ".join(text.split())


def parse_price(value):
    text = clean(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def display_price(value):
    if value is None:
        return ""
    if value == value.to_integral_value():
        return str(int(value))
    return format(value.normalize(), "f")


def multiplied(value, multiplier):
    parsed = parse_price(value)
    if parsed is None:
        return ""
    return display_price(parsed * multiplier)


def normalized_key(sku):
    return clean(sku).lower()


def price_for_sku(sku, price_map):
    key = normalized_key(sku)
    if key in price_map:
        return price_map[key]
    match = re.match(r"^(\d+)\s*pcs\s*(.+)$", clean(sku), flags=re.IGNORECASE)
    if not match:
        return ["", "", ""]
    multiplier = int(match.group(1))
    base_key = normalized_key(match.group(2))
    if base_key not in price_map:
        return ["", "", ""]
    base = price_map[base_key]
    return [multiplied(v, multiplier) for v in base]


def load_prices(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    header_row = None
    headers = []
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        labels = [clean(v) for v in row]
        if "商品 SKU" in labels or "SKU" in labels:
            header_row = idx
            headers = labels
            break
    if header_row is None:
        raise SystemExit("价格表里没有找到 SKU 表头")
    sku_idx = headers.index("商品 SKU") if "商品 SKU" in headers else headers.index("SKU")
    daily_idx = headers.index("日常售价")
    activity_idx = headers.index("平台活动价")
    original_idx = headers.index("上架原价")
    out = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        sku = clean(row[sku_idx] if sku_idx < len(row) else "")
        if not sku:
            continue
        out[normalized_key(sku)] = [
            clean(row[original_idx] if original_idx < len(row) else ""),
            clean(row[daily_idx] if daily_idx < len(row) else ""),
            clean(row[activity_idx] if activity_idx < len(row) else ""),
        ]
    return out


payload = json.loads(DATA.read_text(encoding="utf-8"))
records = sorted(payload["records"], key=lambda r: (r["store"], r["sku"], r["source_row"]))
price_map = load_prices(PRICE_XLSX)

info = run_cli(["sheets", "+info", "--as", "user", "--spreadsheet-token", TOKEN])
sheet_ids = {s["title"]: s["sheet_id"] for s in info["data"]["sheets"]["sheets"]}

# Reset all SKU statuses to normal and clear notes/date.
status_rows = sorted(payload["summary"], key=lambda x: x["sku"])
write_range(sheet_ids["SKU状态总控"], "B2", [["正常"] for _ in status_rows], chunk_size=50)
write_range(sheet_ids["SKU状态总控"], "E2", [["", ""] for _ in status_rows], chunk_size=50)

# Employee dispatch list keeps all rows, with prices visible but status normal.
dispatch_status = [["正常", "", "", "", ""] for _ in records]
write_range(sheet_ids["员工执行清单"], "C2", dispatch_status, chunk_size=50)
dispatch_prices = [["上架原价", "日常售价", "平台活动价"]]
for r in records:
    dispatch_prices.append(price_for_sku(r["sku"], price_map))
write_range(sheet_ids["员工执行清单"], "H1", dispatch_prices, chunk_size=20)

# Store tabs: show prices next to each SKU, with status normal.
records_by_store = defaultdict(list)
for r in records:
    records_by_store[r["store"]].append(r)
matched = 0
missing = 0
for store, rows in records_by_store.items():
    sid = sheet_ids.get(f"店铺-{store}"[:100])
    if not sid:
        continue
    rows = sorted(rows, key=lambda x: (x["sku"], x["source_row"]))
    write_range(sid, "B2", [["正常", "", "", ""] for _ in rows], chunk_size=50)
    matrix = [["上架原价", "日常售价", "平台活动价"]]
    for r in rows:
        prices = price_for_sku(r["sku"], price_map)
        if any(prices):
            matched += 1
        else:
            missing += 1
        matrix.append(prices)
    write_range(sid, "F1", matrix, chunk_size=20)

# Pending dispatch should only show actual pending changes; after reset it is empty.
pending = [
    ["当前待处理任务", "0 条", "", "", "", "", "", "", "", ""],
    ["使用方式", "只有下架、开预购、关预购、待确认的 SKU 会出现在这里；正常商品不显示，避免干扰。", "", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", "", "", ""],
    ["店铺", "SKU", "需要修改为", "员工操作", "完成状态", "更新时间", "备注", "上架原价", "日常售价", "平台活动价"],
]
write_range(sheet_ids["待处理派工"], "A1", [[""] * 10 for _ in range(200)], chunk_size=50)
write_range(sheet_ids["待处理派工"], "A1", pending, chunk_size=20)

print(
    json.dumps(
        {
            "price_skus": len(price_map),
            "store_price_rows_matched": matched,
            "store_price_rows_missing": missing,
            "status": "all_normal",
            "pending_rows": 0,
        },
        ensure_ascii=False,
    )
)
