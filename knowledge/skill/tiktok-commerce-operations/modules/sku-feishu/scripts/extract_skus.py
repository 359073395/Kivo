import json
import os
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(os.environ["SHOPDIR"])
OUT = Path(os.environ["OUT_JSON"])


def shop_name(path: Path) -> str:
    name = path.stem
    for suffix in ("_ID", "-ID", "_id", "-id"):
        if name.endswith(suffix):
            return name[: -len(suffix)]
    return name


def clean_sku(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return " ".join(text.split())


records = []
exceptions = []
file_stats = []

for path in sorted(BASE.glob("*.xlsx")):
    if path.name.startswith("~$"):
        continue

    store = shop_name(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        ws.reset_dimensions()

        headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
        sku_col_idx = None
        for idx, header in enumerate(headers):
            if str(header or "").strip().lower() == "sku":
                sku_col_idx = idx
                break
        if sku_col_idx is None:
            exceptions.append(
                {
                    "store": store,
                    "file": path.name,
                    "row": "",
                    "issue": "未找到 sku 列",
                    "value": "",
                }
            )
            file_stats.append({"store": store, "file": path.name, "rows": 0, "unique_skus": 0})
            continue

        row_count = 0
        store_skus = []
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            row_count += 1
            sku = clean_sku(row[sku_col_idx] if sku_col_idx < len(row) else None)
            if not sku:
                exceptions.append(
                    {
                        "store": store,
                        "file": path.name,
                        "row": row_no,
                        "issue": "SKU 为空",
                        "value": "",
                    }
                )
                continue
            records.append(
                {
                    "sku": sku,
                    "store": store,
                    "source_file": path.name,
                    "source_row": row_no,
                }
            )
            store_skus.append(sku)

        counts = Counter(store_skus)
        for sku, count in counts.items():
            if count > 1:
                exceptions.append(
                    {
                        "store": store,
                        "file": path.name,
                        "row": "",
                        "issue": "同一店铺内重复 SKU",
                        "value": f"{sku} x {count}",
                    }
                )
        file_stats.append(
            {
                "store": store,
                "file": path.name,
                "rows": row_count,
                "unique_skus": len(counts),
            }
        )
    except Exception as exc:
        exceptions.append(
            {
                "store": store,
                "file": path.name,
                "row": "",
                "issue": "读取文件失败",
                "value": repr(exc),
            }
        )
        file_stats.append({"store": store, "file": path.name, "rows": 0, "unique_skus": 0})

stores_by_sku = defaultdict(set)
rows_by_sku = defaultdict(int)
for record in records:
    stores_by_sku[record["sku"]].add(record["store"])
    rows_by_sku[record["sku"]] += 1

summary = []
for sku in sorted(stores_by_sku):
    stores = sorted(stores_by_sku[sku])
    summary.append(
        {
            "sku": sku,
            "store_count": len(stores),
            "stores": "、".join(stores),
            "row_count": rows_by_sku[sku],
        }
    )

OUT.write_text(
    json.dumps(
        {
            "records": records,
            "summary": summary,
            "exceptions": exceptions,
            "file_stats": file_stats,
        },
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)

print(
    json.dumps(
        {
            "files": len(file_stats),
            "records": len(records),
            "unique_skus": len(summary),
            "exceptions": len(exceptions),
        },
        ensure_ascii=False,
    )
)
