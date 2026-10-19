"""Reprice the verified SKU universe from exact source rows; retain mapping snapshots."""
import argparse
import hashlib
import json
import os
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_CEILING
from pathlib import Path

from openpyxl import load_workbook

RULE = 'cat_e_plus_8000_then_5pct_min_y_20260922'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reprice(vault):
    skill = vault / 'skill/tiktok-commerce-operations'
    pointer_path = skill / 'references/current-warehouse-source.json'
    pointer = json.loads(pointer_path.read_text(encoding='utf-8'))
    config = json.loads((vault / '辅助/配置/machine-paths.local.json').read_text(encoding='utf-8'))
    root = Path(config['tiktokSkuRoot'])
    src = root / pointer['relativePath']
    ysrc = root / pointer['price2']['relativePath']
    target = root / pointer['derivedUnifiedPrices']['relativePath']
    assert pointer['pricingRule']['id'] == RULE
    assert sha(src) == pointer['sha256']
    assert sha(target) == pointer['derivedUnifiedPrices']['sha256']
    before_hash = sha(target)
    wb = load_workbook(target)
    cached = load_workbook(target, read_only=True, data_only=True)
    cached['SKU价格表'].reset_dimensions()
    baseline = list(cached['SKU价格表'].values)
    assert wb['数据来源'].cell(5, 5).value == sha(ysrc), 'Y source changed; full source reconciliation required'
    aw = load_workbook(src, read_only=True, data_only=True)
    yw = load_workbook(ysrc, read_only=True, data_only=True)
    aw.active.reset_dimensions(); yw['售价活动价'].reset_dimensions()
    ar = list(aw.active.values); yr = list(yw['售价活动价'].values)
    amap = defaultdict(set); ymap = defaultdict(set)
    for r in ar[1:]:
        if len(r) > 4 and r[1]: amap[str(r[1]).strip()].add(str(r[4]))
    for r in yr[2:]:
        if len(r) > 2 and r[0]: ymap[str(r[0]).strip()].add(str(r[2]))
    refs = skill / 'modules/sku-feishu/references'
    aliases = []
    for name in ['confirmed-warehouse-sku-aliases.json', 'confirmed-y-sku-aliases.json']:
        aliases.append({a['storeSku']: a['warehouseSku'] for a in json.loads((refs/name).read_text(encoding='utf-8'))['aliases']})
    prices = {}; changed = []
    ws = wb['SKU价格表']
    ws['A2'] = '价格1 = ROUNDUP((9.12仓库E列猫店价格 + 8000) × 1.05, 0)；价格2 = Y仓平台活动价；最终取低。Npcs基础拆算按基础价×数量。'
    ws['I4'] = '总加价(先+8000再加5%)'
    for n in range(5, ws.max_row + 1):
        r = list(baseline[n-1])
        sku, base, qty = r[:3]
        if not sku: continue
        assert sku not in prices
        # Existing unresolved SKUs remain unresolved, never filled from guesses.
        if r[16] != '可用':
            for c in range(1,21): ws.cell(n,c).value = r[c-1]
            prices[sku] = r
            continue
        p1 = None
        if r[5] is not None:
            source = ar[int(r[5])-1]
            key = sku if r[3] != 'Npcs基础SKU×数量' else base
            assert source[1] == r[4]
            assert key == source[1] or aliases[0].get(key) == source[1]
            assert len(amap[source[1]]) == 1
            multiplier = int(qty) if r[3] == 'Npcs基础SKU×数量' else 1
            if source[4] is not None:
                cat = Decimal(str(source[4]))
                assert cat > 0 and Decimal(str(r[7])) == cat * multiplier
                unit_price = ((cat + 8000) * Decimal('1.05')).quantize(Decimal(1), rounding=ROUND_CEILING)
                p1 = unit_price * multiplier
                extra = p1 - cat * multiplier
                r[8], r[9] = int(extra), int(p1)
            else:
                assert sku == '626-Expce Herbal Tea' and r[9] is None
        source = yr[int(r[12])-1]
        source_key = str(source[0]).strip()
        key = base if r[10] == 'Npcs基础SKU×数量' else sku
        assert source_key == r[11]
        if key != source_key and aliases[1].get(key) != source_key:
            r[13:18] = [None, None, None, '阻断', f'Y仓SKU未精确匹配且无确认别名：{key} / {source_key}']
            for c in range(1,21): ws.cell(n,c).value = r[c-1]
            prices[sku] = r
            continue
        assert len(ymap[source_key]) == 1
        multiplier = int(qty) if r[10] == 'Npcs基础SKU×数量' else 1
        p2 = Decimal(str(source[2])) * multiplier
        assert p2 > 0 and p2 == Decimal(str(r[13]))
        if p1 is None:
            assert sku == '626-Expce Herbal Tea' and p2 == 59520
            final, origin = p2, '价格2（用户确认）'
        else:
            final = min(p1, p2)
            origin = '价格1' if p1 <= p2 else '价格2'
        if final != Decimal(str(r[14])):
            changed.append({'sku':sku,'before':r[14],'after':int(final)})
        r[14], r[15] = int(final), origin
        for c in range(1,21): ws.cell(n,c).value = r[c-1]
        prices[sku] = r
    assert prices['EELHOE Ginger hair'][9] == 27300
    assert prices['EELHOE Ginger hair'][14] == 27300
    assert prices['626-Expce Herbal Tea'][14] == 59520
    assert prices['3pcsGardenia Hand Cream'][9] == 51450*3
    assert prices['100ML GINSENG Light'][14] == 45380  # Y price wins.

    mapping = wb['ProductID-SKU映射']; groups = defaultdict(dict)
    mapping['A1'] = '9月14日店铺映射快照｜价格已按2026-09-22新公式重算；非最新在线清单'
    for n in range(4,mapping.max_row+1):
        folder,pid,sku = mapping.cell(n,1).value,mapping.cell(n,3).value,mapping.cell(n,4).value
        if not sku: continue
        p = prices[sku]
        for col,value in {5:p[9],6:p[13],7:p[14],8:p[15],9:p[16],10:p[17]}.items(): mapping.cell(n,col).value=value
        groups[(folder,pid)][sku] = p
    def card_status(group):
        if len(group)>1:return '多SKU跳过',None
        p=next(iter(group.values()))
        return ('可用',p[14]) if p[16]=='可用' else ('SKU价格阻断',None)
    for n in range(4,mapping.max_row+1):
        key=(mapping.cell(n,1).value,mapping.cell(n,3).value)
        if key in groups:mapping.cell(n,11).value=card_status(groups[key])[0]
    cards = wb['产品卡校验'];cards['A1']='9月14日产品映射快照｜同一Product ID多个SKU整项跳过｜2026-09-22重算'
    for n in range(4,cards.max_row+1):
        key=(cards.cell(n,1).value,cards.cell(n,3).value)
        if key not in groups:continue
        group=groups[key];status,value=card_status(group)
        cards.cell(n,6).value=' | '.join(str(p[14]) if p[14] is not None else '阻断' for p in group.values())
        cards.cell(n,7).value=value;cards.cell(n,8).value=status
    anomalies = wb['异常清单']
    for n in range(4,anomalies.max_row+1):
        if anomalies.cell(n,1).value == '多SKU异价阻断':
            folder=anomalies.cell(n,2).value;pid=str(anomalies.cell(n,3).value).split(' / ')[-1]
            assert (folder,pid) in groups
            anomalies.cell(n,1).value='多SKU跳过'
            anomalies.cell(n,6).value=' | '.join(str(p[14]) for p in groups[(folder,pid)].values())
            anomalies.cell(n,7).value='用户2026-09-22确认：同一Product ID多个SKU整项跳过，无需重复询问'
    anomalies['A2']='保留9月14日映射快照异常；价格按2026-09-22重算，多SKU商品按用户规则跳过；其他源冲突继续阻断。'
    for sku,p in prices.items():
        if p[16]=='阻断' and sku != '3cpsKISS ROSE SPF':
            anomalies.append(['SKU价格阻断','本次精确源校验',None,sku,None,None,p[17]])
    rules=wb['规则说明']
    rules['B5']='价格1=ROUNDUP((当前9.12仓库E列猫店价格+8000)×1.05,0) IDR；先加8000再加5%，无旧3000封顶。E列已含普店+3000，不重复加3000。2026-09-22用户确认。'
    rules['B11']='同一Product ID多个不同SKU整项跳过，不论价格是否相同；单SKU缺价或冲突仍阻断。'
    rules['B12']='SKU价格按2026-09-22公式重算；Product ID映射仍为9月14日历史快照。实际Flash必须使用各店当前重新导出的商品表。'
    rules['B19']=sum(card_status(g)[0]=='可用' for g in groups.values())
    rules['B15']=len(prices)
    rules['B18']=len(groups)
    pid_shops=defaultdict(set)
    for folder,pid in groups:pid_shops[pid].add(folder)
    rules['B20']=sum(len(folders)>1 for folders in pid_shops.values())
    rules['B16']=sum(p[16]=='可用' for p in prices.values())
    rules['B17']=sum(p[16]!='可用' for p in prices.values())
    # Replace old worked example, including old values in columns C onwards.
    for n in range(22,26):
        for cell in rules[n]:
            if cell.__class__.__name__ != 'MergedCell':cell.value=None
    rules['A22']='新公式核验示例';rules['B22']='EELHOE Ginger hair'
    rules['A23']='(仓库E列+8000)×1.05';rules['B23']=27300
    rules['A24']='Y仓活动价';rules['B24']=36620
    rules['A25']='最终取低';rules['B25']=27300
    sources=wb['数据来源']
    sources.cell(4,3).value=str(src);sources.cell(5,3).value=str(ysrc)
    for n in range(6,23):sources.cell(n,1).value='历史产品表快照（2026-09-14，非当前在线证明）'
    sources.append(['当前定价规则',None,'2026-09-22用户：猫店价格+8000，然后还要再加5%，再与Y仓活动价取低',datetime.now().isoformat(),RULE])
    temp=target.with_name(target.stem+'.codex-tmp.xlsx')
    assert not temp.exists()
    try:
        wb.save(temp)
        check=load_workbook(temp,read_only=True,data_only=True)
        s=check['SKU价格表'];s.reset_dimensions();got=list(s.values)[4:]
        assert len(got)==len(prices)
        for row in got:
            expected=prices[row[0]]
            assert tuple(row[8:18])==tuple(expected[8:18])
        check.close();cached.close();wb.close();aw.close();yw.close()
        assert sha(target)==before_hash
        os.replace(temp,target)
    finally:
        if temp.exists():temp.unlink()
    pointer['derivedUnifiedPrices'].update(status='current_for_source',sha256=sha(target),pricingRule=RULE,
        lastKnownGenerationDate='2026-09-22',price2Sha256=sha(ysrc),
        verification=f'{len(prices)} SKU rows source-checked; (E+8000)*1.05 verified; mappings remain 2026-09-14 snapshots')
    pointer_path.write_text(json.dumps(pointer,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return {'pricing_rule':RULE,'sku_count':len(prices),'changed_final_prices':len(changed),'changes':changed,'unified_sha256':sha(target)}


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--vault-root',type=Path,required=True);p.add_argument('--report',type=Path,required=True)
    args=p.parse_args();result=reprice(args.vault_root)
    args.report.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k!='changes'},ensure_ascii=False,indent=2))
