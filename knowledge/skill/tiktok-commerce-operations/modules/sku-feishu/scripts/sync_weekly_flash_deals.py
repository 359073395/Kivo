from __future__ import annotations

import hashlib
import json
import os
import re
import zipfile
from collections import Counter, defaultdict
from copy import deepcopy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from npcs_unit_rule import validate_npcs_records

from flash_price_rules import (
    apply_confirmed_y_sku_aliases,
    apply_y_platform_activity_fallback,
)


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS = {"m": MAIN_NS}


def required_path(name: str) -> Path:
    value = os.environ.get(name, "").strip()
    if not value:
        raise RuntimeError(f"{name} is required")
    path = Path(value).resolve()
    if not path.exists():
        raise FileNotFoundError(f"{name} does not exist: {path}")
    return path


PRODUCT_ROOT = required_path("PRODUCT_ROOT")
STORE_MAPPING = required_path("STORE_MAPPING")
UNIFIED_PRICE_WORKBOOK = required_path("UNIFIED_PRICE_WORKBOOK")
PRICE_WORKBOOK = (
    Path(os.environ["PRICE_WORKBOOK"]).resolve()
    if os.environ.get("PRICE_WORKBOOK", "").strip()
    else None
)
Y_PRICE_WORKBOOK = (
    Path(os.environ["Y_PRICE_WORKBOOK"]).resolve()
    if os.environ.get("Y_PRICE_WORKBOOK", "").strip()
    else None
)
Y_ALIAS_FILE = (
    Path(__file__).resolve().parents[1]
    / "references"
    / "confirmed-y-sku-aliases.json"
)
CONFIRMED_FLASH_EXCLUSIONS_FILE = (
    Path(__file__).resolve().parents[1]
    / "references"
    / "confirmed-flash-product-exclusions.json"
)
Y_ALIAS_PRICE_WORKBOOK = (
    Path(os.environ["Y_ALIAS_PRICE_WORKBOOK"]).resolve()
    if os.environ.get("Y_ALIAS_PRICE_WORKBOOK", "").strip()
    else None
)
CURRENT_DISCOUNT_RATE = (
    Decimal(os.environ["CURRENT_DISCOUNT_RATE"].strip())
    if os.environ.get("CURRENT_DISCOUNT_RATE", "").strip()
    else None
)
APPLY = os.environ.get("APPLY", "").strip() == "1"
EXCLUDE_NPCS = os.environ.get("EXCLUDE_NPCS", "").strip() == "1"


def normalize(value: object) -> str:
    return " ".join(str(value or "").strip().casefold().split())


def decimal_value(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def product_id_text(value: object) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float):
        return format(value, ".0f")
    return str(value).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_npcs(value: object) -> bool:
    return bool(re.match(r"^\s*\d+\s*pcs", str(value or ""), re.IGNORECASE))


def header_index(headers: list[str], name: str) -> int:
    try:
        return headers.index(name)
    except ValueError as exc:
        raise KeyError(f"Missing required column: {name}") from exc


def load_unified_prices() -> tuple[dict[str, dict[str, object]], list[dict[str, object]]]:
    workbook = load_workbook(UNIFIED_PRICE_WORKBOOK, read_only=True, data_only=True)
    if "SKU价格表" not in workbook.sheetnames:
        workbook.close()
        raise KeyError("Unified price workbook has no SKU价格表 sheet")
    worksheet = workbook["SKU价格表"]
    worksheet.reset_dimensions()

    header_row = None
    headers: list[str] = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        current = [str(value or "").strip() for value in row]
        if "SKU" in current and "最终价格" in current and "状态" in current:
            header_row = row_number
            headers = current
            break
    if header_row is None:
        workbook.close()
        raise KeyError("Unified price workbook has no SKU/最终价格/状态 header row")

    sku_col = header_index(headers, "SKU")
    price1_col = header_index(headers, "价格1")
    price2_col = header_index(headers, "价格2")
    final_col = header_index(headers, "最终价格")
    source_col = header_index(headers, "选用来源")
    status_col = header_index(headers, "状态")
    blocker_col = header_index(headers, "阻断原因")

    candidates: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        sku = str(row[sku_col] or "").strip()
        if not sku:
            continue
        candidates[normalize(sku)].append(
            {
                "sku": sku,
                "price1": decimal_value(row[price1_col]),
                "price2": decimal_value(row[price2_col]),
                "selected_price": decimal_value(row[final_col]),
                "selected_source": str(row[source_col] or "").strip(),
                "status": str(row[status_col] or "").strip(),
                "blocker": str(row[blocker_col] or "").strip(),
            }
        )
    workbook.close()

    result: dict[str, dict[str, object]] = {}
    conflicts: list[dict[str, object]] = []
    for key, records in candidates.items():
        signatures = {
            (
                record["price1"],
                record["price2"],
                record["selected_price"],
                record["status"],
                record["blocker"],
            )
            for record in records
        }
        if len(signatures) != 1:
            conflicts.append(
                {
                    "sku": records[0]["sku"],
                    "reason": "unified price workbook contains conflicting duplicate SKU rows",
                }
            )
            continue
        result[key] = records[0]
    conflicts.extend(validate_npcs_records(result.values()))
    return result, conflicts


def load_main_prices() -> tuple[dict[tuple[str, str], dict[str, object]], list[dict[str, object]]]:
    workbook = load_workbook(PRICE_WORKBOOK, read_only=True, data_only=True)
    worksheet = workbook["全部店铺SKU"]
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    shop_col = header_index(headers, "店铺")
    sku_col = header_index(headers, "商品SKU")
    flash_col = header_index(headers, "SKU平时闪购总价（展示价48%）")
    status_col = header_index(headers, "原处理状态")
    source_col = header_index(headers, "源行号")

    candidates: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        shop = str(row[shop_col] or "").strip()
        sku = str(row[sku_col] or "").strip()
        if not shop or not sku:
            continue
        candidates[(normalize(shop), normalize(sku))].append(
            {
                "shop": shop,
                "sku": sku,
                "flash_price": decimal_value(row[flash_col]),
                "status": str(row[status_col] or "").strip(),
                "source_rows": str(row[source_col] or "").strip(),
            }
        )
    workbook.close()

    result: dict[tuple[str, str], dict[str, object]] = {}
    conflicts: list[dict[str, object]] = []
    for key, records in candidates.items():
        signatures = {
            (record["flash_price"], record["status"])
            for record in records
        }
        if len(signatures) != 1:
            conflicts.append(
                {
                    "shop": records[0]["shop"],
                    "sku": records[0]["sku"],
                    "signatures": sorted(str(item) for item in signatures),
                }
            )
            continue
        result[key] = records[0]
    return result, conflicts


def load_y_prices() -> tuple[dict[str, Decimal], list[dict[str, object]]]:
    workbook = load_workbook(Y_PRICE_WORKBOOK, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    sku_col = header_index(headers, "SKU")
    activity_name = next(
        (
            name
            for name in (
                "平台活动价(IDR)",
                "平台活动价",
                "活动价(IDR)",
                "活动价",
            )
            if name in headers
        ),
        None,
    )
    if activity_name is None:
        raise KeyError("Y warehouse workbook has no platform activity price column")
    activity_col = header_index(headers, activity_name)

    candidates: dict[str, set[Decimal]] = defaultdict(set)
    labels: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        sku = str(row[sku_col] or "").strip()
        price = decimal_value(row[activity_col])
        if not sku or price is None:
            continue
        candidates[normalize(sku)].add(price)
        labels[normalize(sku)].add(sku)
    workbook.close()

    result: dict[str, Decimal] = {}
    conflicts: list[dict[str, object]] = []
    for sku, prices in candidates.items():
        if len(prices) == 1:
            result[sku] = next(iter(prices))
        else:
            conflicts.append(
                {
                    "sku": sorted(labels[sku])[0],
                    "prices": sorted(str(price) for price in prices),
                }
            )
    aliases = load_confirmed_y_aliases()
    if not aliases:
        return result, conflicts
    if Y_ALIAS_PRICE_WORKBOOK is None:
        raise RuntimeError(
            "Y_ALIAS_PRICE_WORKBOOK is required because confirmed Y SKU aliases exist"
        )
    if not Y_ALIAS_PRICE_WORKBOOK.is_file():
        raise FileNotFoundError(
            f"Y_ALIAS_PRICE_WORKBOOK does not exist: {Y_ALIAS_PRICE_WORKBOOK}"
        )

    alias_source_candidates = load_y_alias_source_candidates(
        Y_ALIAS_PRICE_WORKBOOK
    )
    alias_source_prices: dict[str, Decimal] = {}
    for target in set(aliases.values()):
        prices = alias_source_candidates.get(target, set())
        if len(prices) != 1:
            conflicts.append(
                {
                    "sku": target,
                    "prices": sorted(str(price) for price in prices),
                    "reason": "confirmed alias target missing or conflicting",
                }
            )
            continue
        alias_source_prices[target] = next(iter(prices))
    if conflicts:
        return result, conflicts
    return apply_confirmed_y_sku_aliases(result, alias_source_prices, aliases), conflicts


def load_confirmed_y_aliases() -> dict[str, str]:
    if not Y_ALIAS_FILE.is_file():
        return {}
    payload = json.loads(Y_ALIAS_FILE.read_text(encoding="utf-8"))
    aliases: dict[str, str] = {}
    for item in payload.get("aliases", []):
        store_sku = normalize(item.get("storeSku"))
        warehouse_sku = normalize(item.get("warehouseSku"))
        if not store_sku or not warehouse_sku:
            raise ValueError(f"Invalid confirmed Y SKU alias entry: {item}")
        existing = aliases.get(store_sku)
        if existing is not None and existing != warehouse_sku:
            raise ValueError(f"Conflicting confirmed Y SKU alias: {item['storeSku']}")
        aliases[store_sku] = warehouse_sku
    return aliases


def load_confirmed_flash_exclusions() -> dict[tuple[str, str], dict[str, object]]:
    if not CONFIRMED_FLASH_EXCLUSIONS_FILE.is_file():
        return {}
    payload = json.loads(
        CONFIRMED_FLASH_EXCLUSIONS_FILE.read_text(encoding="utf-8")
    )
    result: dict[tuple[str, str], dict[str, object]] = {}
    for item in payload.get("exclusions", []):
        folder = str(item.get("folder") or "").strip()
        product_id = product_id_text(item.get("productId"))
        if not folder or not product_id:
            raise ValueError(f"Invalid confirmed Flash exclusion: {item}")
        key = (folder, product_id)
        if key in result:
            raise ValueError(f"Duplicate confirmed Flash exclusion: {key}")
        result[key] = item
    return result


def load_y_alias_source_candidates(path: Path) -> dict[str, set[Decimal]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = (
        workbook["售价活动价"]
        if "售价活动价" in workbook.sheetnames
        else workbook.active
    )
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers: list[str] | None = None
    for row_number, row in enumerate(rows, 1):
        current = [str(value or "").strip() for value in row]
        if "商品 SKU" in current and "平台活动价" in current:
            headers = current
            break
        if row_number >= 20:
            break
    if headers is None:
        workbook.close()
        raise KeyError("Y alias price workbook has no 商品 SKU / 平台活动价 header row")

    sku_col = header_index(headers, "商品 SKU")
    activity_col = header_index(headers, "平台活动价")
    candidates: dict[str, set[Decimal]] = defaultdict(set)
    for row in rows:
        sku = str(row[sku_col] or "").strip()
        price = decimal_value(row[activity_col])
        if sku and price is not None:
            candidates[normalize(sku)].add(price)
    workbook.close()
    return candidates


def load_products(path: Path) -> tuple[dict[str, list[dict[str, object]]], list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = (
        workbook["tiktok_chanpin_"]
        if "tiktok_chanpin_" in workbook.sheetnames
        else workbook.active
    )
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    sku_col = header_index(headers, "sku")
    product_col = header_index(headers, "产品id")
    listed_col = header_index(headers, "价格(站点币种)")
    local_col = header_index(headers, "本地展示价")
    title_col = header_index(headers, "产品标题")

    products: dict[str, list[dict[str, object]]] = defaultdict(list)
    missing_ids: list[dict[str, object]] = []
    for row_number, row in enumerate(rows, 2):
        sku = str(row[sku_col] or "").strip()
        if not sku:
            continue
        product_id = product_id_text(row[product_col])
        listed_price = decimal_value(row[listed_col])
        if listed_price is None:
            listed_price = decimal_value(row[local_col])
        record = {
            "row": row_number,
            "sku": sku,
            "product_id": product_id,
            "listed_price": listed_price,
            "title": str(row[title_col] or "").strip(),
        }
        if not product_id:
            missing_ids.append(record)
        else:
            products[product_id].append(record)
    workbook.close()
    return products, missing_ids


def load_existing_flash(path: Path) -> dict[str, Decimal]:
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
        if broken is not None:
            raise ValueError(f"Corrupt Flash Deal workbook: {path} ({broken})")
        required = {"xl/sharedStrings.xml", "xl/worksheets/sheet1.xml"}
        if not required.issubset(set(archive.namelist())):
            raise ValueError(f"Flash Deal workbook is missing required OOXML: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    first = normalize(headers[0] if headers else "").replace("_", " ")
    second = normalize(headers[1] if len(headers) > 1 else "")
    price_header_ok = (
        ("deal" in second and "price" in second)
        or ("harga" in second and "penawaran" in second)
    )
    if "product" not in first or "id" not in first or not price_header_ok:
        raise ValueError(f"Unexpected Flash Deal headers in {path}: {headers[:2]}")

    result: dict[str, Decimal] = {}
    for row_number, row in enumerate(rows, 2):
        product_id = product_id_text(row[0] if row else None)
        price = decimal_value(row[1] if len(row) > 1 else None)
        if not product_id and price is None:
            continue
        if not product_id or price is None:
            raise ValueError(f"Partial Flash Deal row in {path}: row {row_number}")
        if product_id in result and result[product_id] != price:
            raise ValueError(f"Conflicting duplicate product ID in {path}: {product_id}")
        result[product_id] = price
    workbook.close()
    return result


def analyze_store(
    products: dict[str, list[dict[str, object]]],
    unified_prices: dict[str, dict[str, object]],
) -> tuple[list[dict[str, object]], dict[str, list[dict[str, object]]]]:
    reasons = {
        "multi_sku": [],
        "npcs": [],
        "clearance": [],
        "multi_price": [],
        "missing_price": [],
        "missing_display": [],
        "unsafe_price": [],
        "y_platform_fallback_blocker": [],
    }
    eligible: list[dict[str, object]] = []

    for product_id, variants in sorted(products.items()):
        skus = sorted({str(item["sku"]) for item in variants})
        normalized_skus = [normalize(sku) for sku in skus]
        base = {"product_id": product_id, "skus": skus}
        # User confirmed on 2026-09-22: skip the entire Product ID when
        # it maps to more than one distinct SKU, even if prices are equal.
        if len(skus) > 1:
            reasons["multi_sku"].append(
                {**base, "reason": "用户规则：同一Product ID多个SKU，整项跳过"}
            )
            continue
        if EXCLUDE_NPCS and any(is_npcs(sku) for sku in skus):
            reasons["npcs"].append(
                {**base, "reason": "当前批次按用户指令临时排除Npcs"}
            )
            continue
        price_records = [unified_prices.get(sku) for sku in normalized_skus]
        blocked = [
            record
            for record in price_records
            if record is None
            or record["status"] != "可用"
            or record["selected_price"] is None
        ]
        if blocked:
            blockers = sorted(
                {
                    str(record["blocker"] or "统一SKU价格缺失")
                    if record is not None
                    else "SKU未进入统一价格表"
                    for record in blocked
                }
            )
            target = "clearance" if any(
                any(marker in blocker for marker in ("清仓", "下架"))
                for blocker in blockers
            ) else "missing_price"
            reasons[target].append({**base, "reason": "；".join(blockers)})
            continue

        prices = {
            record["selected_price"]
            for record in price_records
            if record is not None and record["selected_price"] is not None
        }
        if len(prices) != 1:
            reasons["multi_price"].append(
                {
                    **base,
                    "reason": "同一Product ID的多个SKU最终价格不同",
                    "prices": sorted(str(price) for price in prices),
                }
            )
            continue

        candidate = next(iter(prices))
        sources = sorted(
            {
                str(record["selected_source"] or "统一价格表")
                for record in price_records
                if record is not None
            }
        )
        offer_price = int(candidate.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        eligible.append(
            {
                "product_id": product_id,
                "offer_price": offer_price,
                "skus": skus,
                "price_source": " + ".join(sources),
            }
        )

    return eligible, reasons


def patch_flash(path: Path, rows: list[dict[str, object]]) -> dict[str, object]:
    temporary = path.with_name(f"{path.name}.codex-tmp")
    if temporary.exists():
        raise FileExistsError(f"Temporary workbook already exists: {temporary}")

    with zipfile.ZipFile(path, "r") as source:
        original_entries = source.namelist()
        shared_root = etree.fromstring(source.read("xl/sharedStrings.xml"))
        sheet_root = etree.fromstring(source.read("xl/worksheets/sheet1.xml"))
        sheet_data = sheet_root.find("m:sheetData", NS)
        if sheet_data is None:
            raise ValueError(f"Flash Deal workbook has no sheetData: {path}")
        sheet_rows = {
            int(row.get("r")): row
            for row in sheet_data.findall("m:row", NS)
        }
        style_row = sheet_rows.get(2)
        if style_row is None:
            raise ValueError(f"Flash Deal workbook has no style row 2: {path}")
        style_row_template = deepcopy(style_row)

        def column_name(cell: etree._Element) -> str:
            match = re.match(r"[A-Z]+", str(cell.get("r")))
            if match is None:
                raise ValueError(f"Invalid cell reference in {path}")
            return match.group(0)

        def cell_map(row: etree._Element) -> dict[str, etree._Element]:
            return {column_name(cell): cell for cell in row.findall("m:c", NS)}

        styles = cell_map(style_row)
        if "A" not in styles or "B" not in styles:
            raise ValueError(f"Flash Deal style row lacks A/B cells: {path}")
        product_style = str(styles["A"].get("s") or "0")
        price_style = str(styles["B"].get("s") or "0")

        def ensure_row(row_number: int) -> etree._Element:
            row = sheet_rows.get(row_number)
            if row is not None:
                return row
            # Always clone the pristine blank style row.  Cloning `style_row`
            # after row 2 has been populated duplicates its Product_id and
            # Deal Price cells in every appended row.
            row = deepcopy(style_row_template)
            row.set("r", str(row_number))
            for cell in list(row.findall("m:c", NS)):
                cell.set("r", f"{column_name(cell)}{row_number}")
                if column_name(cell) in {"A", "B"}:
                    row.remove(cell)
            sheet_data.append(row)
            sheet_rows[row_number] = row
            return row

        def remove_ab(row: etree._Element) -> None:
            for cell in list(row.findall("m:c", NS)):
                if column_name(cell) in {"A", "B"}:
                    row.remove(cell)

        def add_cell(row: etree._Element, column: str, style: str) -> etree._Element:
            cell = etree.Element(
                f"{{{MAIN_NS}}}c",
                r=f"{column}{row.get('r')}",
                s=style,
            )
            existing = row.findall("m:c", NS)
            rank = ord(column[0])
            for index, current in enumerate(existing):
                if ord(column_name(current)[0]) > rank:
                    row.insert(index, cell)
                    return cell
            row.append(cell)
            return cell

        for row_number, row in sheet_rows.items():
            if row_number >= 2:
                remove_ab(row)

        base_shared = len(shared_root.findall("m:si", NS))
        for offset, item in enumerate(rows):
            row_number = offset + 2
            row = ensure_row(row_number)
            product_cell = add_cell(row, "A", product_style)
            product_cell.set("t", "s")
            shared_item = etree.SubElement(shared_root, f"{{{MAIN_NS}}}si")
            shared_text = etree.SubElement(shared_item, f"{{{MAIN_NS}}}t")
            shared_text.text = str(item["product_id"])
            product_value = etree.SubElement(product_cell, f"{{{MAIN_NS}}}v")
            product_value.text = str(base_shared + offset)

            price_cell = add_cell(row, "B", price_style)
            price_value = etree.SubElement(price_cell, f"{{{MAIN_NS}}}v")
            price_value.text = str(int(item["offer_price"]))

        shared_total = base_shared + len(rows)
        shared_root.set("count", str(shared_total))
        shared_root.set("uniqueCount", str(shared_total))

        dimension = sheet_root.find("m:dimension", NS)
        if dimension is not None:
            max_existing = max(sheet_rows) if sheet_rows else 1
            dimension.set("ref", f"A1:N{max(12, max_existing, len(rows) + 1)}")

        shared_xml = etree.tostring(
            shared_root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )
        sheet_xml = etree.tostring(
            sheet_root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )

        with zipfile.ZipFile(temporary, "w") as output:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "xl/sharedStrings.xml":
                    content = shared_xml
                elif entry.filename == "xl/worksheets/sheet1.xml":
                    content = sheet_xml
                output.writestr(entry, content)

    try:
        with zipfile.ZipFile(temporary, "r") as check:
            if check.namelist() != original_entries:
                raise ValueError(f"OOXML entry list changed: {path}")
            broken = check.testzip()
            if broken is not None:
                raise ValueError(f"Generated Flash Deal is corrupt: {path} ({broken})")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()

    return {"ooxml_entries": len(original_entries)}


def verify_flash(path: Path, rows: list[dict[str, object]]) -> None:
    expected_ids = [str(item["product_id"]) for item in rows]
    if len(expected_ids) != len(set(expected_ids)):
        duplicates = sorted(
            product_id
            for product_id, count in Counter(expected_ids).items()
            if count > 1
        )
        raise ValueError(
            f"Duplicate expected product IDs for {path}: {duplicates[:5]}"
        )

    with zipfile.ZipFile(path, "r") as archive:
        sheet_root = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    sheet_data = sheet_root.find("m:sheetData", NS)
    if sheet_data is None:
        raise ValueError(f"Flash Deal workbook has no sheetData: {path}")

    def cell_column(cell: etree._Element) -> str:
        match = re.match(r"[A-Z]+", str(cell.get("r")))
        if match is None:
            raise ValueError(f"Invalid cell reference in {path}")
        return match.group(0)

    def has_value(cell: etree._Element) -> bool:
        value = cell.find("m:v", NS)
        if value is not None and str(value.text or "").strip():
            return True
        inline = cell.find("m:is", NS)
        return inline is not None and bool("".join(inline.itertext()).strip())

    populated_rows = 0
    for row in sheet_data.findall("m:row", NS):
        row_number = int(row.get("r"))
        if row_number < 2:
            continue
        cells = row.findall("m:c", NS)
        refs = [str(cell.get("r")) for cell in cells]
        duplicate_refs = sorted(
            ref for ref, count in Counter(refs).items() if count > 1
        )
        if duplicate_refs:
            raise ValueError(
                f"Duplicate cell references in {path}: "
                f"row {row_number} {duplicate_refs}"
            )
        by_column = {cell_column(cell): cell for cell in cells}
        a_value = has_value(by_column["A"]) if "A" in by_column else False
        b_value = has_value(by_column["B"]) if "B" in by_column else False
        if a_value or b_value:
            populated_rows += 1
            if not a_value or not b_value:
                raise ValueError(
                    f"Partial Flash Deal row in {path}: row {row_number}"
                )
            if sum(1 for cell in cells if cell_column(cell) == "A") != 1:
                raise ValueError(f"Invalid Product_id cell count in {path}: row {row_number}")
            if sum(1 for cell in cells if cell_column(cell) == "B") != 1:
                raise ValueError(f"Invalid Deal Price cell count in {path}: row {row_number}")
            for column in ("C", "D"):
                if column in by_column and has_value(by_column[column]):
                    raise ValueError(
                        f"Unexpected value in column {column} of {path}: row {row_number}"
                    )
    if populated_rows != len(rows):
        raise ValueError(
            f"Flash Deal structural row mismatch for {path}: "
            f"expected={len(rows)} actual={populated_rows}"
        )

    actual = load_existing_flash(path)
    expected = {
        str(item["product_id"]): Decimal(str(item["offer_price"]))
        for item in rows
    }
    if actual != expected:
        missing = sorted(set(expected) - set(actual))
        extra = sorted(set(actual) - set(expected))
        changed = sorted(
            key for key in set(actual) & set(expected)
            if actual[key] != expected[key]
        )
        raise ValueError(
            f"Flash Deal readback mismatch for {path}: "
            f"missing={missing[:5]} extra={extra[:5]} changed={changed[:5]}"
        )


def main() -> None:
    mapping = json.loads(STORE_MAPPING.read_text(encoding="utf-8"))
    expected_folders = {entry["folder"] for entry in mapping}
    actual_folders = {path.name for path in PRODUCT_ROOT.iterdir() if path.is_dir()}
    if expected_folders != actual_folders:
        raise ValueError(
            "Store folder set changed: "
            f"missing={sorted(expected_folders - actual_folders)} "
            f"unexpected={sorted(actual_folders - expected_folders)}"
        )

    unified_prices, unified_conflicts = load_unified_prices()
    if unified_conflicts:
        raise ValueError(
            f"Unified price workbook has conflicting SKU rows: {unified_conflicts[:5]}"
        )

    prepared: list[dict[str, object]] = []
    for entry in mapping:
        folder = PRODUCT_ROOT / entry["folder"]
        product_path = folder / entry["product_file"]
        flash_path = folder / "Flash Deal.xlsx"
        if not product_path.is_file() or not flash_path.is_file():
            raise FileNotFoundError(f"Missing mapped input for {entry['folder']}")

        products, missing_ids = load_products(product_path)
        old_flash = load_existing_flash(flash_path)
        eligible, reasons = analyze_store(
            products,
            unified_prices,
        )
        eligible_ids = {str(item["product_id"]) for item in eligible}
        if not eligible_ids.issubset(set(products)):
            raise ValueError(f"Eligible IDs are not current online IDs: {entry['folder']}")
        prepared.append(
            {
                "folder": entry["folder"],
                "product_path": product_path,
                "flash_path": flash_path,
                "product_ids": len(products),
                "product_detail_rows": sum(len(items) for items in products.values()),
                "missing_product_ids": len(missing_ids),
                "old_flash": old_flash,
                "eligible": eligible,
                "content_changed": {
                    str(item["product_id"]): Decimal(str(item["offer_price"]))
                    for item in eligible
                } != old_flash,
                "reasons": reasons,
                "before_hash": sha256_file(flash_path),
            }
        )

    confirmed_exclusions = load_confirmed_flash_exclusions()
    all_price_blockers = [
        {
            "folder": store["folder"],
            **item,
        }
        for store in prepared
        for reason_name in ("missing_price", "multi_price")
        for item in store["reasons"][reason_name]
    ]
    confirmed_price_blockers = [
        item
        for item in all_price_blockers
        if (str(item["folder"]), str(item["product_id"])) in confirmed_exclusions
    ]
    price_blockers = [
        item
        for item in all_price_blockers
        if (str(item["folder"]), str(item["product_id"])) not in confirmed_exclusions
    ]
    if APPLY and price_blockers:
        raise ValueError(
            "Unified pricing has product-level blockers; ask the user before writing: "
            f"{price_blockers[:10]}"
        )

    if APPLY:
        for store in prepared:
            if store["content_changed"]:
                patch_flash(store["flash_path"], store["eligible"])
        for store in prepared:
            verify_flash(store["flash_path"], store["eligible"])

    report = []
    for store in prepared:
        eligible = store["eligible"]
        old_flash = store["old_flash"]
        old_ids = set(old_flash)
        new_ids = {str(item["product_id"]) for item in eligible}
        exclusions = {
            key: len(items) for key, items in store["reasons"].items()
        }
        report.append(
            {
                "folder": store["folder"],
                "online_product_ids": store["product_ids"],
                "product_detail_rows": store["product_detail_rows"],
                "missing_product_ids": store["missing_product_ids"],
                "old_flash_rows": len(old_flash),
                "new_flash_rows": len(eligible),
                "added": len(new_ids - old_ids),
                "removed": len(old_ids - new_ids),
                "price_sources": dict(
                    Counter(str(item["price_source"]) for item in eligible)
                ),
                "exclusions": exclusions,
                "price_blocker_details": {
                    "multi_price": store["reasons"]["multi_price"],
                    "missing_price": store["reasons"]["missing_price"],
                },
                "before_sha256": store["before_hash"],
                "after_sha256": (
                    sha256_file(store["flash_path"]) if APPLY else None
                ),
                "readback_verified": APPLY,
                "content_changed": store["content_changed"],
            }
        )

    print(
        json.dumps(
            {
                "apply": APPLY,
                "stores": len(report),
                "unified_price_workbook": str(UNIFIED_PRICE_WORKBOOK),
                "pricing_rule": "默认不参加大促；价格1与价格2取低",
                "exclude_npcs_for_this_run": EXCLUDE_NPCS,
                "confirmed_price_blockers_excluded": len(confirmed_price_blockers),
                "unconfirmed_price_blockers": len(price_blockers),
                "report": report,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
