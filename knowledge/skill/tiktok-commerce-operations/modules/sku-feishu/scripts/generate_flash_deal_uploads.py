from __future__ import annotations

import json
import os
import re
import zipfile
from collections import Counter, defaultdict
from copy import deepcopy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from lxml import etree


ROOT_TEXT = os.environ.get("TIKTOK_SKU_ROOT", "").strip()
if not ROOT_TEXT:
    raise RuntimeError(
        "TIKTOK_SKU_ROOT is required. Load the current machine path before running; "
        "the script will not fall back to the Skill directory or a legacy drive."
    )
ROOT = Path(ROOT_TEXT)
SETTINGS_FILE = ROOT / "settings.json"
SETTINGS: dict[str, object] = {}
if SETTINGS_FILE.is_file():
    SETTINGS = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))

SHOP_DIRECTORY = Path(
    os.environ.get("SHOP_DIRECTORY", ROOT / "01_店铺上架产品")
)
PRICE_WORKBOOK = Path(os.environ["PRICE_WORKBOOK"])
FLASH_DEAL_TEMPLATE_DIRECTORY_TEXT = str(
    os.environ.get("ZHANFU_BROWSER_ROOT")
    or os.environ.get("FLASH_DEAL_TEMPLATE_DIRECTORY")
    or SETTINGS.get("zhanfu_browser_root")
    or ""
).strip()
FLASH_DEAL_TEMPLATE_DIRECTORY = (
    Path(FLASH_DEAL_TEMPLATE_DIRECTORY_TEXT)
    if FLASH_DEAL_TEMPLATE_DIRECTORY_TEXT
    else None
)
OUTPUT_DIRECTORY_TEXT = os.environ.get("OUTPUT_DIRECTORY", "").strip()
OUTPUT_DIRECTORY = Path(OUTPUT_DIRECTORY_TEXT) if OUTPUT_DIRECTORY_TEXT else None
AUDIT_WORKBOOK_TEXT = os.environ.get("AUDIT_WORKBOOK", "").strip()
AUDIT_WORKBOOK = Path(AUDIT_WORKBOOK_TEXT) if AUDIT_WORKBOOK_TEXT else None
CURRENT_DISCOUNT_RATE_TEXT = str(
    os.environ.get("CURRENT_DISCOUNT_RATE")
    or SETTINGS.get("current_discount_rate")
    or ""
).strip()
if not CURRENT_DISCOUNT_RATE_TEXT:
    raise RuntimeError(
        "CURRENT_DISCOUNT_RATE is required. Do not reuse a historical discount "
        "rate without confirming the current pricing source."
    )
CURRENT_DISCOUNT_RATE = Decimal(CURRENT_DISCOUNT_RATE_TEXT)
TARGET_STORES = {
    " ".join(item.strip().casefold().split())
    for item in os.environ.get("TARGET_STORES", "").split(",")
    if item.strip()
}


def normalize(value: object) -> str:
    return " ".join(str(value or "").strip().casefold().split())


def is_npcs_sku(value: object) -> bool:
    return bool(re.match(r"^\s*\d+\s*pcs", str(value or ""), re.IGNORECASE))


def decimal_value(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def product_id_text(value: object) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".0f")
    return str(value).strip()


def safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", value).strip().rstrip(".")
    return cleaned or "store"


def resolve_store_template(shop: str) -> Path | None:
    if FLASH_DEAL_TEMPLATE_DIRECTORY is None:
        return None
    candidates = [
        FLASH_DEAL_TEMPLATE_DIRECTORY / shop / "Flash Deal.xlsx",
        FLASH_DEAL_TEMPLATE_DIRECTORY / f"{shop}.xlsx",
    ]
    return next((path for path in candidates if path.is_file()), None)


def header_index(headers: list[str], name: str) -> int:
    try:
        return headers.index(name)
    except ValueError as exc:
        raise KeyError(f"Missing required column: {name}") from exc


def load_prices() -> dict[tuple[str, str], dict[str, object]]:
    workbook = load_workbook(PRICE_WORKBOOK, read_only=True, data_only=True)
    worksheet = workbook["全部店铺SKU"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]

    shop_col = header_index(headers, "店铺")
    sku_col = header_index(headers, "商品SKU")
    flash_col = header_index(headers, "SKU平时闪购总价（展示价48%）")
    display_col = header_index(headers, "最终应展示价格")
    status_col = header_index(headers, "原处理状态")
    source_rows_col = header_index(headers, "源行号")

    result: dict[tuple[str, str], dict[str, object]] = {}
    for row in rows:
        shop = str(row[shop_col] or "").strip()
        sku = str(row[sku_col] or "").strip()
        if not shop or not sku:
            continue
        result[(shop, normalize(sku))] = {
            "shop": shop,
            "sku": sku,
            "flash_price": decimal_value(row[flash_col]),
            "display_price": decimal_value(row[display_col]),
            "status": str(row[status_col] or "").strip(),
            "source_rows": str(row[source_rows_col] or "").strip(),
        }

    workbook.close()
    return result


def load_store_products(path: Path) -> tuple[dict[str, list[dict[str, object]]], list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]

    sku_col = header_index(headers, "sku")
    product_id_col = header_index(headers, "产品id")
    listed_price_col = header_index(headers, "价格(站点币种)")
    local_display_col = header_index(headers, "本地展示价")
    title_col = header_index(headers, "产品标题")

    products: dict[str, list[dict[str, object]]] = defaultdict(list)
    missing_product_ids: list[dict[str, object]] = []

    for row_number, row in enumerate(rows, 2):
        sku = str(row[sku_col] or "").strip()
        if not sku:
            continue
        product_id = product_id_text(row[product_id_col])
        listed_price = decimal_value(row[listed_price_col])
        if listed_price is None:
            listed_price = decimal_value(row[local_display_col])
        record = {
            "row": row_number,
            "sku": sku,
            "product_id": product_id,
            "title": row[title_col],
            "listed_price": listed_price,
            "current_actual_price": (
                listed_price * CURRENT_DISCOUNT_RATE
                if listed_price is not None
                else None
            ),
        }
        if not product_id:
            missing_product_ids.append(record)
            continue
        products[product_id].append(record)

    workbook.close()
    return products, missing_product_ids


def analyze_store(
    shop: str,
    products: dict[str, list[dict[str, object]]],
    price_map: dict[tuple[str, str], dict[str, object]],
) -> dict[str, list[dict[str, object]]]:
    result: dict[str, list[dict[str, object]]] = {
        "eligible": [],
        "conflicts": [],
        "npcs_excluded": [],
        "clearance": [],
        "pending": [],
        "selling_price_risk": [],
    }

    for product_id, variants in sorted(products.items()):
        variant_details: list[dict[str, object]] = []
        for variant in variants:
            price = price_map.get((shop, normalize(variant["sku"])))
            variant_details.append({**variant, "price": price})

        skus = sorted({str(item["sku"]) for item in variant_details})
        clearance_items = [
            item
            for item in variant_details
            if item["price"]
            and any(
                marker in str(item["price"]["status"])
                for marker in ("清仓", "下架")
            )
        ]
        if clearance_items:
            result["clearance"].append(
                {
                    "shop": shop,
                    "product_id": product_id,
                    "skus": skus,
                    "reason": "产品ID包含清仓/下架SKU",
                    "affected_skus": sorted(
                        {str(item["sku"]) for item in clearance_items}
                    ),
                }
            )
            continue

        npcs_items = [
            item for item in variant_details if is_npcs_sku(item["sku"])
        ]
        if npcs_items:
            result["npcs_excluded"].append(
                {
                    "shop": shop,
                    "product_id": product_id,
                    "skus": skus,
                    "affected_skus": sorted(
                        {str(item["sku"]) for item in npcs_items}
                    ),
                    "reason": "批量Flash Deal模板对同一产品ID的全部规格使用同一价格；Npcs组合装必须按规格单独定价",
                }
            )
            continue

        unavailable = [
            item
            for item in variant_details
            if item["price"] is None
            or item["price"]["flash_price"] is None
            or item["price"]["status"] != "已生成"
        ]
        if unavailable:
            result["pending"].append(
                {
                    "shop": shop,
                    "product_id": product_id,
                    "skus": skus,
                    "reason": "产品ID内存在未匹配、缺价或待确认SKU",
                    "affected_skus": sorted(
                        {str(item["sku"]) for item in unavailable}
                    ),
                }
            )
            continue

        price_options: dict[Decimal, list[str]] = defaultdict(list)
        for item in variant_details:
            flash_price = item["price"]["flash_price"]
            assert isinstance(flash_price, Decimal)
            price_options[flash_price].append(str(item["sku"]))

        if len(price_options) != 1:
            result["conflicts"].append(
                {
                    "shop": shop,
                    "product_id": product_id,
                    "skus": skus,
                    "reason": "同一产品ID包含多个不同闪购价",
                    "price_options": "; ".join(
                        f"{price:,.0f}: {', '.join(sorted(set(option_skus)))}"
                        for price, option_skus in sorted(price_options.items())
                    ),
                }
            )
            continue

        offer_decimal = next(iter(price_options))
        offer_price = int(offer_decimal.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        current_actual_prices = [
            item["current_actual_price"]
            for item in variant_details
            if item["current_actual_price"] is not None
        ]
        listed_prices = [
            item["listed_price"]
            for item in variant_details
            if item["listed_price"] is not None
        ]
        if not current_actual_prices:
            result["pending"].append(
                {
                    "shop": shop,
                    "product_id": product_id,
                    "skus": skus,
                    "reason": "缺少当前展示价，无法校验闪购价必须低于售价",
                    "affected_skus": skus,
                }
            )
            continue

        minimum_current_actual = min(current_actual_prices)
        if Decimal(offer_price) >= minimum_current_actual:
            result["selling_price_risk"].append(
                {
                    "shop": shop,
                    "product_id": product_id,
                    "skus": skus,
                    "offer_price": offer_price,
                    "minimum_current_actual": minimum_current_actual,
                    "reason": "闪购价不低于当前5折实际售价/L30参考线",
                }
            )
            continue

        result["eligible"].append(
            {
                "shop": shop,
                "product_id": product_id,
                "skus": skus,
                "offer_price": offer_price,
                "minimum_current_actual": minimum_current_actual,
                "minimum_listed_price": min(listed_prices) if listed_prices else None,
            }
        )

    return result


def write_upload(
    shop: str,
    rows: list[dict[str, object]],
    template_path: Path,
) -> Path:
    if OUTPUT_DIRECTORY is None:
        output_path = template_path
    else:
        OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
        output_path = OUTPUT_DIRECTORY / f"Flash Deal_不参加大促_{safe_filename(shop)}.xlsx"
    temporary_output = output_path.with_name(f"{output_path.name}.codex-tmp")
    if temporary_output.exists():
        raise FileExistsError(f"Temporary output already exists: {temporary_output}")
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns = {"m": main_ns}

    with zipfile.ZipFile(template_path, "r") as source_zip:
        original_entries = source_zip.namelist()
        required_entries = {
            "xl/sharedStrings.xml",
            "xl/worksheets/sheet1.xml",
        }
        missing_entries = required_entries - set(source_zip.namelist())
        if missing_entries:
            raise ValueError(
                f"Store template is missing required OOXML entries: "
                f"{template_path} ({', '.join(sorted(missing_entries))})"
            )
        shared_root = etree.fromstring(source_zip.read("xl/sharedStrings.xml"))
        sheet_root = etree.fromstring(source_zip.read("xl/worksheets/sheet1.xml"))
        sheet_data = sheet_root.find("m:sheetData", ns)
        if sheet_data is None:
            raise ValueError("Official Flash Deal template has no sheetData")

        sheet_rows = {
            int(row.get("r")): row for row in sheet_data.findall("m:row", ns)
        }
        template_row = sheet_rows[max(sheet_rows)]
        style_row = sheet_rows.get(2)
        if style_row is None:
            raise ValueError(f"Store template has no data row 2: {template_path}")

        def ensure_row(row_number: int):
            existing = sheet_rows.get(row_number)
            if existing is not None:
                return existing
            new_row = deepcopy(template_row)
            new_row.set("r", str(row_number))
            new_row.set("spans", "1:14")
            for cell in new_row.findall("m:c", ns):
                column = re.match(r"[A-Z]+", str(cell.get("r"))).group(0)
                cell.set("r", f"{column}{row_number}")
            sheet_data.append(new_row)
            sheet_rows[row_number] = new_row
            return new_row

        def cell_map(row):
            return {
                re.match(r"[A-Z]+", str(cell.get("r"))).group(0): cell
                for cell in row.findall("m:c", ns)
            }

        style_cells = cell_map(style_row)
        if "A" not in style_cells or "B" not in style_cells:
            raise ValueError(
                f"Store template row 2 has no Product ID/Deal Price cells: {template_path}"
            )
        product_style = str(style_cells["A"].get("s") or "0")
        price_style = str(style_cells["B"].get("s") or "0")

        def ensure_cell(row, column: str, style: str):
            cells = cell_map(row)
            if column in cells:
                cell = cells[column]
                cell.set("s", style)
                return cell
            cell = etree.Element(f"{{{main_ns}}}c", r=f"{column}{row.get('r')}", s=style)
            order = ord(column) - ord("A")
            inserted = False
            for index, current in enumerate(row.findall("m:c", ns)):
                current_column = re.match(r"[A-Z]+", str(current.get("r"))).group(0)
                if ord(current_column[0]) - ord("A") > order:
                    row.insert(index, cell)
                    inserted = True
                    break
            if not inserted:
                row.append(cell)
            return cell

        base_shared_count = len(shared_root.findall("m:si", ns))
        for row_number, item in enumerate(rows, 2):
            row = ensure_row(row_number)
            row.set("spans", "1:14")

            product_id = str(item["product_id"])
            shared_item = etree.SubElement(shared_root, f"{{{main_ns}}}si")
            shared_text = etree.SubElement(shared_item, f"{{{main_ns}}}t")
            shared_text.text = product_id
            shared_index = base_shared_count + row_number - 2

            product_cell = ensure_cell(row, "A", product_style)
            for child in list(product_cell):
                product_cell.remove(child)
            product_cell.set("t", "s")
            product_value = etree.SubElement(product_cell, f"{{{main_ns}}}v")
            product_value.text = str(shared_index)

            price_cell = ensure_cell(row, "B", price_style)
            for child in list(price_cell):
                price_cell.remove(child)
            price_cell.attrib.pop("t", None)
            price_value = etree.SubElement(price_cell, f"{{{main_ns}}}v")
            price_value.text = str(int(item["offer_price"]))

        shared_total = base_shared_count + len(rows)
        shared_root.set("count", str(shared_total))
        shared_root.set("uniqueCount", str(shared_total))

        dimension = sheet_root.find("m:dimension", ns)
        if dimension is not None:
            dimension.set("ref", f"A1:N{max(12, len(rows) + 1)}")

        shared_xml = etree.tostring(
            shared_root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )
        sheet_xml = etree.tostring(
            sheet_root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )

        with zipfile.ZipFile(temporary_output, "w") as output_zip:
            for entry in source_zip.infolist():
                content = source_zip.read(entry.filename)
                if entry.filename == "xl/sharedStrings.xml":
                    content = shared_xml
                elif entry.filename == "xl/worksheets/sheet1.xml":
                    content = sheet_xml
                output_zip.writestr(entry, content)

    try:
        with zipfile.ZipFile(temporary_output, "r") as check_zip:
            if check_zip.namelist() != original_entries:
                raise ValueError(
                    f"OOXML entry list changed unexpectedly: {template_path}"
                )
            broken_entry = check_zip.testzip()
            if broken_entry is not None:
                raise ValueError(
                    f"Generated workbook contains a corrupt ZIP entry: {broken_entry}"
                )
        os.replace(temporary_output, output_path)
    finally:
        if temporary_output.exists():
            temporary_output.unlink()
    return output_path


def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    dark = "1F4E78"
    thin = Side(style="thin", color="D9E1F2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_number in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column_number)
        width = max(
            len(str(worksheet.cell(row_number, column_number).value or ""))
            for row_number in range(1, min(worksheet.max_row, 200) + 1)
        )
        worksheet.column_dimensions[letter].width = min(max(width + 2, 12), 50)


def write_audit(
    summaries: list[dict[str, object]],
    all_results: dict[str, list[dict[str, object]]],
    missing_product_ids: list[dict[str, object]],
) -> None:
    if AUDIT_WORKBOOK is None:
        return
    AUDIT_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_rows = [
        [
            item["shop"],
            item["source_rows"],
            item["unique_skus"],
            item["product_ids"],
            item["eligible"],
            item["conflicts"],
            item["npcs_excluded"],
            item["clearance"],
            item["pending"],
            item["selling_price_risk"],
            item["missing_product_id"],
        ]
        for item in summaries
    ]
    add_sheet(
        workbook,
        "汇总",
        [
            "店铺",
            "源产品行",
            "唯一SKU",
            "产品ID数",
            "可上传行",
            "产品级价格冲突",
            "Npcs批量排除",
            "清仓下架排除",
            "缺价待确认",
            "当前售价/L30风险",
            "缺产品ID",
        ],
        summary_rows,
    )

    add_sheet(
        workbook,
        "Npcs批量排除",
        ["店铺", "产品ID", "关联SKU", "Npcs SKU", "原因"],
        [
            [
                item["shop"],
                item["product_id"],
                ", ".join(item["skus"]),
                ", ".join(item["affected_skus"]),
                item["reason"],
            ]
            for item in all_results["npcs_excluded"]
        ],
    )

    add_sheet(
        workbook,
        "可上传明细",
        ["店铺", "产品ID", "关联SKU", "闪购价", "当前5折最低售价", "最低展示价"],
        [
            [
                item["shop"],
                item["product_id"],
                ", ".join(item["skus"]),
                item["offer_price"],
                item["minimum_current_actual"],
                item["minimum_listed_price"],
            ]
            for item in all_results["eligible"]
        ],
    )
    add_sheet(
        workbook,
        "产品级价格冲突",
        ["店铺", "产品ID", "关联SKU", "不同价格", "原因"],
        [
            [
                item["shop"],
                item["product_id"],
                ", ".join(item["skus"]),
                item["price_options"],
                item["reason"],
            ]
            for item in all_results["conflicts"]
        ],
    )
    add_sheet(
        workbook,
        "清仓下架排除",
        ["店铺", "产品ID", "关联SKU", "清仓下架SKU", "原因"],
        [
            [
                item["shop"],
                item["product_id"],
                ", ".join(item["skus"]),
                ", ".join(item["affected_skus"]),
                item["reason"],
            ]
            for item in all_results["clearance"]
        ],
    )
    add_sheet(
        workbook,
        "缺价待确认",
        ["店铺", "产品ID", "关联SKU", "受影响SKU", "原因"],
        [
            [
                item["shop"],
                item["product_id"],
                ", ".join(item["skus"]),
                ", ".join(item["affected_skus"]),
                item["reason"],
            ]
            for item in all_results["pending"]
        ],
    )
    add_sheet(
        workbook,
        "售价L30风险",
        ["店铺", "产品ID", "关联SKU", "建议闪购价", "当前5折最低售价", "原因"],
        [
            [
                item["shop"],
                item["product_id"],
                ", ".join(item["skus"]),
                item["offer_price"],
                item["minimum_current_actual"],
                item["reason"],
            ]
            for item in all_results["selling_price_risk"]
        ],
    )
    add_sheet(
        workbook,
        "缺产品ID",
        ["店铺", "源行", "SKU", "标题"],
        [
            [item["shop"], item["row"], item["sku"], item["title"]]
            for item in missing_product_ids
        ],
    )

    workbook.save(AUDIT_WORKBOOK)
    workbook.close()


def main() -> None:
    price_map = load_prices()
    summaries: list[dict[str, object]] = []
    all_results: dict[str, list[dict[str, object]]] = {
        "eligible": [],
        "conflicts": [],
        "npcs_excluded": [],
        "clearance": [],
        "pending": [],
        "selling_price_risk": [],
    }
    all_missing_product_ids: list[dict[str, object]] = []
    generated_files: list[str] = []
    prepared_uploads: list[tuple[str, list[dict[str, object]], Path]] = []

    if (
        FLASH_DEAL_TEMPLATE_DIRECTORY is None
        or not FLASH_DEAL_TEMPLATE_DIRECTORY.is_dir()
    ):
        raise FileNotFoundError(
            "Zhanfu browser download root was not found: "
            f"{FLASH_DEAL_TEMPLATE_DIRECTORY}. Set zhanfu_browser_root in "
            f"{SETTINGS_FILE} or set ZHANFU_BROWSER_ROOT."
        )

    store_paths = [
        path
        for path in sorted(SHOP_DIRECTORY.glob("*.xlsx"))
        if not path.name.startswith("~$")
    ]
    if TARGET_STORES:
        store_paths = [
            path for path in store_paths if normalize(path.stem) in TARGET_STORES
        ]
        found_stores = {normalize(path.stem) for path in store_paths}
        missing_store_exports = sorted(TARGET_STORES - found_stores)
        if missing_store_exports:
            raise FileNotFoundError(
                "Requested store exports were not found: "
                + ", ".join(missing_store_exports)
            )
    store_templates: dict[str, Path] = {}
    missing_templates: list[str] = []
    for path in store_paths:
        shop = path.stem
        template_path = resolve_store_template(shop)
        if template_path is None:
            missing_templates.append(shop)
        else:
            store_templates[shop] = template_path

    if missing_templates:
        expected = FLASH_DEAL_TEMPLATE_DIRECTORY / "<store>" / "Flash Deal.xlsx"
        raise FileNotFoundError(
            "Missing freshly downloaded store-specific Flash Deal templates for: "
            + ", ".join(missing_templates)
            + f". Expected layout: {expected}"
        )

    for path in store_paths:
        shop = path.stem
        products, missing_product_ids = load_store_products(path)
        for item in missing_product_ids:
            item["shop"] = shop
        result = analyze_store(shop, products, price_map)
        prepared_uploads.append(
            (shop, result["eligible"], store_templates[shop])
        )

        raw_rows = sum(len(variants) for variants in products.values()) + len(missing_product_ids)
        unique_skus = len(
            {
                normalize(item["sku"])
                for variants in products.values()
                for item in variants
            }
            | {normalize(item["sku"]) for item in missing_product_ids}
        )
        summaries.append(
            {
                "shop": shop,
                "source_rows": raw_rows,
                "unique_skus": unique_skus,
                "product_ids": len(products),
                "eligible": len(result["eligible"]),
                "conflicts": len(result["conflicts"]),
                "npcs_excluded": len(result["npcs_excluded"]),
                "clearance": len(result["clearance"]),
                "pending": len(result["pending"]),
                "selling_price_risk": len(result["selling_price_risk"]),
                "missing_product_id": len(missing_product_ids),
            }
        )
        for key in all_results:
            all_results[key].extend(result[key])
        all_missing_product_ids.extend(missing_product_ids)

    if AUDIT_WORKBOOK is not None:
        write_audit(summaries, all_results, all_missing_product_ids)

    # All stores are analyzed before the first official workbook is changed.
    for shop, rows, template_path in prepared_uploads:
        upload_path = write_upload(shop, rows, template_path)
        generated_files.append(str(upload_path))

    print(
        json.dumps(
            {
                "stores": len(summaries),
                "upload_files": len(generated_files),
                "eligible_rows": len(all_results["eligible"]),
                "conflicts": len(all_results["conflicts"]),
                "npcs_excluded": len(all_results["npcs_excluded"]),
                "clearance": len(all_results["clearance"]),
                "pending": len(all_results["pending"]),
                "selling_price_risk": len(all_results["selling_price_risk"]),
                "missing_product_id": len(all_missing_product_ids),
                "audit_workbook": (
                    str(AUDIT_WORKBOOK) if AUDIT_WORKBOOK is not None else None
                ),
                "output_directory": (
                    str(OUTPUT_DIRECTORY)
                    if OUTPUT_DIRECTORY is not None
                    else "original store-specific Flash Deal.xlsx updated in place"
                ),
                "template_directory": str(FLASH_DEAL_TEMPLATE_DIRECTORY),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
